#!/usr/bin/env python3
"""
所属 Skill：Skill 3 — 猎头触达文案生成与批量导出 Skill
前置依赖：Python 3.11+、openpyxl>=3.1。
Agent 调用入口：
    python recruiter_message_generator_script.py \
      --input outreach-input.json --output-dir deliverables

边界：本脚本只按标准化画像中的 outreach_directive 渲染文案和 Excel。
它不联网、不评估候选人、不决定语言/力度/是否触达，也不执行自动发送。
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - dependency failure path
    raise SystemExit("Missing dependency: install openpyxl>=3.1") from exc


SCHEMA_VERSION = "1.0"
PRIVATE_EMAIL_DOMAINS = {
    "gmail.com", "googlemail.com", "outlook.com", "hotmail.com", "live.com",
    "yahoo.com", "icloud.com", "me.com", "proton.me", "protonmail.com",
    "qq.com", "163.com", "126.com", "sina.com", "foxmail.com",
}
VALID_MODES = {"formal", "admin_forward", "archive_only", "hold"}
VALID_LOCALES = {"zh_cn", "en_us", "en_eu"}
VALID_INTENSITIES = {"direct", "balanced", "low_pressure"}
ULTRA_FAST_EXPORT_PROFILE = "ultra_fast_single_sheet"


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def compact(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def unique_strings(values: Iterable[Any] | Any | None) -> list[str]:
    if values is None:
        values = []
    elif isinstance(values, (str, bytes)):
        values = [values]
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        cleaned = compact(value)
        key = cleaned.casefold()
        if cleaned and key not in seen:
            seen.add(key)
            result.append(cleaned)
    return result


def get_metric(candidate: dict[str, Any], key: str) -> Any:
    metric = ((candidate.get("profile") or {}).get("academic_metrics") or {}).get(key)
    if isinstance(metric, dict):
        return metric.get("value")
    return metric


def safe_cell(value: Any) -> str | int | float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (dict, list)):
        text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    else:
        text = str(value)
    text = text[:32767]
    # Prevent formula injection when the workbook is opened.
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def first_value(values: Iterable[Any], fallback: str = "") -> str:
    for value in values:
        cleaned = compact(value)
        if cleaned:
            return cleaned
    return fallback


def candidate_name(candidate: dict[str, Any]) -> str:
    profile = candidate.get("profile") or {}
    return compact(profile.get("name") or candidate.get("name"))


def directive(candidate: dict[str, Any]) -> dict[str, Any]:
    return candidate.get("outreach_directive") or {}


def candidate_anchor(candidate: dict[str, Any]) -> dict[str, Any]:
    """Apply the Agent-selected anchor preference without making a new decision."""
    profile = candidate.get("profile") or {}
    policy = directive(candidate)
    preferred_url = compact(policy.get("anchor_source_url"))
    preferred_type = compact(policy.get("anchor_preference"))
    anchors = [
        item for item in as_list(profile.get("outreach_anchors"))
        if isinstance(item, dict)
    ]
    if preferred_url:
        for item in anchors:
            if compact(item.get("source_url")) == preferred_url:
                return item
    if preferred_type:
        for item in anchors:
            if compact(item.get("type")) == preferred_type:
                return item
    return profile.get("outreach_anchor") or (anchors[0] if anchors else {})


def contact_is_sendable(contact: dict[str, Any]) -> bool:
    email = compact(contact.get("email")).casefold()
    if "@" not in email:
        return False
    domain = email.rsplit("@", 1)[1]
    return bool(
        domain not in PRIVATE_EMAIL_DOMAINS
        and contact.get("publicly_listed") is True
        and compact(contact.get("source_url"))
    )


def sendable_contacts(candidate: dict[str, Any], contact_type: str) -> list[dict[str, Any]]:
    contacts = ((candidate.get("profile") or {}).get("contacts") or {}).get(contact_type)
    return [
        item for item in as_list(contacts)
        if isinstance(item, dict) and contact_is_sendable(item)
    ]


def validate_controller(payload: dict[str, Any]) -> list[str]:
    controller = payload.get("controller") or {}
    missing = [
        key for key in (
            "organization",
            "contact_email",
            "legal_basis",
            "retention_period",
            "privacy_notice_url",
        )
        if not compact(controller.get(key))
    ]
    return [f"controller_missing:{key}" for key in missing]


def validate_candidate(
    candidate: dict[str, Any],
    suppression: set[str],
    controller_errors: list[str],
) -> list[str]:
    errors: list[str] = []
    name = candidate_name(candidate)
    policy = directive(candidate)
    mode = compact(policy.get("mode"))
    locale = compact(policy.get("locale"))
    intensity = compact(policy.get("intensity"))
    if not name:
        errors.append("missing_candidate_name")
    if mode not in VALID_MODES:
        errors.append("invalid_or_missing_mode")
    if locale not in VALID_LOCALES:
        errors.append("invalid_or_missing_locale")
    if intensity not in VALID_INTENSITIES:
        errors.append("invalid_or_missing_intensity")
    if mode in {"formal", "admin_forward"}:
        errors.extend(controller_errors)
        anchor = candidate_anchor(candidate)
        if not compact(anchor.get("title")):
            errors.append("missing_public_paper_or_talk_anchor")
        if mode == "formal" and not sendable_contacts(candidate, "public_institutional"):
            errors.append("no_sendable_public_institutional_email")
        if mode == "admin_forward" and not sendable_contacts(candidate, "departmental"):
            errors.append("no_sendable_departmental_email")
        score = (candidate.get("jd_match") or {}).get("score")
        if score is not None and float(score) < 40:
            errors.append("jd_score_below_40_requires_archive_only")
        if (candidate.get("mobility") or {}).get("hold_recommended") is True:
            errors.append("mobility_or_crm_hold_requires_hold_mode")
    keys = {
        compact(candidate.get("candidate_key")).casefold(),
        *(compact(item.get("email")).casefold()
          for kind in ("public_institutional", "departmental")
          for item in sendable_contacts(candidate, kind)),
    }
    if any(key and key in suppression for key in keys):
        errors.append("suppression_list_match")
    opportunity = policy.get("opportunity") or {}
    if policy.get("competitor_talent") is True and not as_list(opportunity.get("differentiators")):
        errors.append("competitor_talent_requires_verified_differentiators")
    return unique_strings(errors)


def controller_footer(payload: dict[str, Any], locale: str, source_url: str) -> str:
    controller = payload.get("controller") or {}
    organization = compact(controller.get("organization"))
    contact_email = compact(controller.get("contact_email"))
    privacy_url = compact(controller.get("privacy_notice_url"))
    legal_basis = compact(controller.get("legal_basis"))
    retention_period = compact(controller.get("retention_period"))
    authority_url = compact(controller.get("supervisory_authority_url"))
    source_clause = source_url or "the candidate's institutionally published contact page"
    if locale == "zh_cn":
        text = (
            f"隐私说明：数据控制方为 {organization}。我们仅处理您的姓名、公开职业履历和机构公示邮箱"
            f"（来源：{source_clause}），用于就与您研究方向相关的职业或合作机会与您联系；"
            f"组织核验的处理依据为“{legal_basis}”，计划保留期限为“{retention_period}”。"
            f"您可免费提出异议、退出后续联系，或申请查询、更正、删除相关信息；请回复本邮件或联系 {contact_email}。"
        )
        text += f" 完整隐私说明及投诉渠道：{privacy_url}"
        if authority_url:
            text += f" 监管机构：{authority_url}"
        return text
    if locale == "en_eu":
        text = (
            f"Privacy notice: The data controller is {organization}. We process only your name, public professional "
            f"affiliation, and institutionally published contact details (source: {source_clause}) to contact you about "
            f"a research-related professional or collaboration opportunity. The organisation-verified legal basis is "
            f"“{legal_basis}”, and the stated retention period is “{retention_period}”. You may object or opt out free "
            f"of charge, or request access, rectification, or erasure, by replying or contacting {contact_email}."
        )
        text += f" Full notice and complaint information: {privacy_url}"
        if authority_url:
            text += f" Supervisory authority: {authority_url}"
        return text
    text = (
        f"Privacy notice: {organization} processes only your name, public professional affiliation, and contact "
        f"information published by your institution (source: {source_clause}) to contact you about a relevant "
        f"professional or collaboration opportunity. The organisation-verified legal basis is “{legal_basis}”; "
        f"the stated retention period is “{retention_period}”. To object or opt out free of charge, or request access, "
        f"correction, or deletion, reply or contact {contact_email}."
    )
    text += f" Full privacy notice and complaint information: {privacy_url}"
    if authority_url:
        text += f" Supervisory authority: {authority_url}"
    return text


def opportunity_text(candidate: dict[str, Any], variant: str, locale: str) -> str:
    policy = directive(candidate)
    opportunity = policy.get("opportunity") or {}
    role = compact(opportunity.get("role_title")) or (
        "相关科研岗位" if locale == "zh_cn" else "a relevant research role"
    )
    organization = compact(opportunity.get("organization"))
    platform = compact(opportunity.get("research_platform"))
    team = compact(opportunity.get("team_scope"))
    compensation = compact(opportunity.get("compensation"))
    facts = unique_strings(as_list(opportunity.get("verified_facts")))
    differentiators = unique_strings(as_list(opportunity.get("differentiators")))
    verified = "；".join(facts[:3]) if locale == "zh_cn" else "; ".join(facts[:3])
    diff = "；".join(differentiators[:2]) if locale == "zh_cn" else "; ".join(differentiators[:2])
    if locale == "zh_cn":
        destination = f"{organization}的{role}" if organization else role
        if variant == "academic":
            details = "、".join(value for value in (platform, verified, diff) if value)
            return f"我们正在交流的{destination}与您的研究方向较为贴合" + (
                f"，经核验的机会信息包括：{details}。" if details else "。"
            )
        if variant == "full_time":
            details = "、".join(value for value in (team, compensation, verified, diff) if value)
            return f"该{destination}重视成果落地与团队协作" + (
                f"，经核验的信息包括：{details}。" if details else "。"
            )
        details = "、".join(value for value in (platform, verified) if value)
        return (
            f"即使您目前不考虑全职流动，也想探讨与{destination}开展兼职顾问或联合研究的可能"
            + (f"，合作基础包括：{details}。" if details else "。")
        )
    destination = f"the {role} opportunity at {organization}" if organization else role
    if variant == "academic":
        details = "; ".join(value for value in (platform, verified, diff) if value)
        return f"{destination.capitalize()} appears closely aligned with your research" + (
            f". Verified details include {details}." if details else "."
        )
    if variant == "full_time":
        details = "; ".join(value for value in (team, compensation, verified, diff) if value)
        return f"{destination.capitalize()} emphasizes translation and team-based delivery" + (
            f". Verified details include {details}." if details else "."
        )
    details = "; ".join(value for value in (platform, verified) if value)
    return (
        f"Even if a full-time move is not relevant, there may be scope to explore an advisory or research collaboration "
        f"connected with {destination}" + (f". The verified basis is {details}." if details else ".")
    )


def cta(locale: str, intensity: str, variant: str) -> str:
    if locale == "zh_cn":
        if intensity == "direct":
            return "若方向契合，想邀请您本周或下周用 15 分钟做一次非正式交流；如不便，也完全理解。"
        if intensity == "balanced":
            return "若您愿意，我可以先发送岗位与平台的完整资料，再由您决定是否安排一次简短交流。"
        return "若这个方向恰好与您近期规划相关，您方便时回复即可；若不考虑，也无需专门说明。"
    if locale == "en_eu":
        if intensity == "direct":
            return "If this is relevant, may I propose a brief, non-committal 15-minute conversation next week? I would, of course, understand if the timing is not suitable."
        if intensity == "balanced":
            return "If useful, I would be pleased to send the full role and platform brief first, leaving you to decide whether a short conversation would be worthwhile."
        return "If the subject happens to be relevant to your current plans, please respond at your convenience; there is no expectation to engage."
    if intensity == "direct":
        return "If this is relevant, would you be open to a low-key 15-minute conversation next week? No pressure if the timing is not right."
    if intensity == "balanced":
        return "If helpful, I can send the full role and platform brief first, and you can decide whether a short conversation makes sense."
    return "If this happens to fit your plans, feel free to reply whenever convenient; no response is expected if it does not."


def anchor_opening(candidate: dict[str, Any], locale: str) -> tuple[str, str, str]:
    profile = candidate.get("profile") or {}
    anchor = candidate_anchor(candidate)
    title = compact(anchor.get("title"))
    summary = compact(anchor.get("summary"))
    anchor_type = compact(anchor.get("type"))
    direction = first_value(profile.get("research_directions") or [], "the topic")
    name = candidate_name(candidate)
    if locale == "zh_cn":
        salutation = f"尊敬的{name}老师，"
        if anchor_type == "public_talk":
            opening = f"我近期认真看了您题为《{title}》的公开演讲资料"
        else:
            opening = f"我近期认真阅读了您参与的论文《{title}》"
        opening += f"，其中围绕{summary or direction}的工作让我很受启发。"
        subject_prefix = f"关于《{title}》研究方向的交流"
        return salutation, opening, subject_prefix
    salutation = f"Dear {name},"
    if anchor_type == "public_talk":
        opening = f"I recently reviewed the public materials for your talk, “{title},”"
    else:
        opening = f"I recently read your paper, “{title},”"
    opening += f" and was particularly interested in its work on {summary or direction}."
    subject_prefix = f"Research discussion: {title}"
    return salutation, opening, subject_prefix


def subject(candidate: dict[str, Any], locale: str, variant: str, anchor_subject: str) -> str:
    opportunity = directive(candidate).get("opportunity") or {}
    role = compact(opportunity.get("role_title"))
    if locale == "zh_cn":
        suffixes = {
            "academic": "与科研平台机会",
            "full_time": "与全职岗位机会",
            "advisory": "与顾问合作可能",
        }
        return f"{anchor_subject}{suffixes[variant]}"
    suffixes = {
        "academic": f" — {role or 'research platform'}",
        "full_time": f" — {role or 'research opportunity'}",
        "advisory": " — an exploratory advisory discussion",
    }
    return anchor_subject + suffixes[variant]


def initial_email(
    payload: dict[str, Any],
    candidate: dict[str, Any],
    variant: str,
    source_url: str,
) -> dict[str, Any]:
    policy = directive(candidate)
    locale = policy["locale"]
    intensity = policy["intensity"]
    salutation, opening, anchor_subject = anchor_opening(candidate, locale)
    opportunity = opportunity_text(candidate, variant, locale)
    invitation = cta(locale, intensity, variant)
    body = "\n\n".join(
        [
            f"{salutation}\n{opening}",
            opportunity,
            invitation,
        ]
    )
    return {
        "subject": subject(candidate, locale, variant, anchor_subject),
        "body": body,
        "privacy_footer": controller_footer(payload, locale, source_url),
        "structure": [
            "paper_or_talk_research_opening",
            "verified_opportunity_match",
            "lightweight_invitation",
        ],
    }


def follow_up(
    payload: dict[str, Any],
    candidate: dict[str, Any],
    source_url: str,
) -> dict[str, Any]:
    policy = directive(candidate)
    locale = policy["locale"]
    profile = candidate.get("profile") or {}
    anchor = candidate_anchor(candidate)
    title = compact(anchor.get("title"))
    direction = first_value(profile.get("research_directions") or [], "相关研究方向")
    name = candidate_name(candidate)
    if locale == "zh_cn":
        paragraphs = [
            f"尊敬的{name}老师，想就上周关于《{title}》的邮件做一次简短跟进。",
            f"我更希望把这次联系理解为一次关于{direction}产业应用与科研平台的交流，而不是仓促的岗位推介。",
            "若您愿意，我可以先发一页简要资料；如目前不便或不考虑，完全无需回复。",
        ]
        subject_line = f"跟进：《{title}》相关研究与应用交流"
    else:
        formal = locale == "en_eu"
        paragraphs = [
            f"Dear {name},\nI wanted to make one brief follow-up to my note last week regarding “{title}.”",
            f"My intention is to open a substantive discussion about the practical application of {direction}, rather than to press an immediate role conversation.",
            (
                "If useful, I would be pleased to send a one-page brief first. "
                "There is no need to reply if the timing or subject is not suitable."
                if formal else
                "If useful, I can send a one-page brief first. No need to reply if the timing or topic is not right."
            ),
        ]
        subject_line = f"Follow-up: research and application discussion on {title}"
    return {
        "send_after_days": 7,
        "subject": subject_line,
        "body": "\n\n".join(paragraphs),
        "privacy_footer": controller_footer(payload, locale, source_url),
    }


def linkedin_copy(candidate: dict[str, Any]) -> dict[str, str]:
    policy = directive(candidate)
    locale = policy["locale"]
    profile = candidate.get("profile") or {}
    anchor = candidate_anchor(candidate)
    title = compact(anchor.get("title"))
    name = candidate_name(candidate)
    direction = first_value(profile.get("research_directions") or [], "")
    if locale == "zh_cn":
        note = f"{name}老师您好，我读到您关于《{title}》的工作，希望就{direction}的科研与应用机会做一次低压力交流。"
        message = (
            f"感谢通过。您在《{title}》中的研究与我们正在关注的{direction}方向很贴合。"
            "若您愿意，我可先发一页经核验的机会资料，再由您决定是否交流。"
        )
    elif locale == "en_eu":
        note = f"Dear {name}, I read your work “{title}” and would value a low-pressure discussion about a relevant research opportunity."
        message = (
            f"Thank you for connecting. Your work in “{title}” is closely relevant to a research opportunity we are mapping. "
            "I would be pleased to send a short verified brief first, with no expectation of a call."
        )
    else:
        note = f"Hi {name} — I read your work “{title}” and would value a low-key conversation about a relevant research opportunity."
        message = (
            f"Thanks for connecting. Your work in “{title}” maps closely to a research opportunity we're exploring. "
            "I can send a short verified brief first—no pressure to schedule a call."
        )
    return {"connection_note": note[:300], "direct_message": message[:1500]}


def admin_forward(
    payload: dict[str, Any],
    candidate: dict[str, Any],
    source_url: str,
) -> dict[str, Any]:
    policy = directive(candidate)
    locale = policy["locale"]
    name = candidate_name(candidate)
    anchor = candidate_anchor(candidate)
    title = compact(anchor.get("title"))
    opportunity = policy.get("opportunity") or {}
    role = compact(opportunity.get("role_title")) or (
        "科研合作机会" if locale == "zh_cn" else "a research opportunity"
    )
    if locale == "zh_cn":
        subject_line = f"烦请转达：关于{name}老师《{title}》研究的交流"
        paragraphs = [
            f"您好，我近期认真阅读了{name}老师参与的论文《{title}》，希望就相关研究方向进行一次专业交流。",
            f"本次联系涉及{role}；因未找到老师公开公示的个人机构邮箱，特礼貌询问贵院系是否方便代为转达。",
            "如不便转达，请直接忽略本邮件；我们不会请求任何非公开联系方式。",
        ]
    else:
        subject_line = f"Kind request to forward: research discussion for {name}"
        paragraphs = [
            f"Dear Department Administrator,\nI recently read {name}'s paper, “{title},” and hope to open a professional discussion about the research.",
            f"The contact concerns {role}. As I could not locate a publicly listed institutional email for the researcher, may I kindly ask whether you would forward this note?",
            "If forwarding is not appropriate, please disregard this message. We are not requesting any non-public contact information.",
        ]
    return {
        "subject": subject_line,
        "body": "\n\n".join(paragraphs),
        "privacy_footer": controller_footer(payload, locale, source_url),
    }


def internal_note(candidate: dict[str, Any], errors: list[str]) -> str:
    policy = directive(candidate)
    mode = compact(policy.get("mode"))
    score = (candidate.get("jd_match") or {}).get("score")
    mobility = candidate.get("mobility") or {}
    reasons = unique_strings(
        as_list(policy.get("hold_reasons"))
        + as_list(mobility.get("signals"))
        + errors
    )
    return (
        f"mode={mode}; jd_score={score}; mobility_window={mobility.get('window')}; "
        f"sendable_message=false; reasons={json.dumps(reasons, ensure_ascii=False)}"
    )


def render_candidate(
    payload: dict[str, Any],
    candidate: dict[str, Any],
    suppression: set[str],
    controller_errors: list[str],
) -> tuple[dict[str, Any], dict[str, Any]]:
    errors = validate_candidate(candidate, suppression, controller_errors)
    policy = directive(candidate)
    mode = compact(policy.get("mode"))
    profile = candidate.get("profile") or {}
    contact_type = "departmental" if mode == "admin_forward" else "public_institutional"
    contacts = sendable_contacts(candidate, contact_type)
    contact = contacts[0] if contacts else {}
    source_url = compact(contact.get("source_url"))
    controller_only_errors = bool(errors) and all(
        item.startswith("controller_missing:") for item in errors
    )
    preview_only = (
        payload.get("draft_preview_without_controller") is True
        and controller_only_errors
        and mode in {"formal", "admin_forward"}
    )
    render_payload = payload
    if preview_only:
        render_payload = dict(payload)
        render_payload["controller"] = {
            "organization": "【待填写数据控制方】",
            "contact_email": "【待填写隐私联系人邮箱】",
            "legal_basis": "【待完成并填写合法利益评估或其他适用依据】",
            "retention_period": "【待填写保留期限】",
            "privacy_notice_url": "【待填写隐私政策链接】",
            "supervisory_authority_url": "",
        }
    result: dict[str, Any] = {
        "candidate_key": candidate.get("candidate_key"),
        "name": candidate_name(candidate),
        "mode": mode,
        "locale": policy.get("locale"),
        "intensity": policy.get("intensity"),
        "recipient": compact(contact.get("email")) or None,
        "sendable": False,
        "preview_only": preview_only,
        "validation_errors": errors,
        "initial_emails": {},
        "follow_up": None,
        "linkedin": None,
        "admin_forward": None,
        "internal_note": None,
    }
    if preview_only and mode == "admin_forward":
        result["admin_forward"] = admin_forward(
            render_payload, candidate, source_url
        )
        result["internal_note"] = internal_note(candidate, errors)
    elif preview_only and mode == "formal":
        result["initial_emails"] = {
            "A_academic_depth": initial_email(
                render_payload, candidate, "academic", source_url
            ),
            "B_full_time_mobility": initial_email(
                render_payload, candidate, "full_time", source_url
            ),
            "C_advisory": initial_email(
                render_payload, candidate, "advisory", source_url
            ),
        }
        result["follow_up"] = follow_up(render_payload, candidate, source_url)
        result["linkedin"] = linkedin_copy(candidate)
        result["internal_note"] = internal_note(candidate, errors)
    elif errors or mode in {"archive_only", "hold"}:
        result["internal_note"] = internal_note(candidate, errors)
    elif mode == "admin_forward":
        result["admin_forward"] = admin_forward(payload, candidate, source_url)
        result["sendable"] = True
    elif mode == "formal":
        result["initial_emails"] = {
            "A_academic_depth": initial_email(
                payload, candidate, "academic", source_url
            ),
            "B_full_time_mobility": initial_email(
                payload, candidate, "full_time", source_url
            ),
            "C_advisory": initial_email(
                payload, candidate, "advisory", source_url
            ),
        }
        result["follow_up"] = follow_up(payload, candidate, source_url)
        result["linkedin"] = linkedin_copy(candidate)
        result["sendable"] = True
    audit = {
        "candidate_key": candidate.get("candidate_key"),
        "name": candidate_name(candidate),
        "requested_mode": mode,
        "sendable": result["sendable"],
        "preview_only": preview_only,
        "recipient_source_url": source_url or None,
        "recipient_is_public_institutional": bool(contact),
        "gdpr_footer_added": result["sendable"],
        "placeholder_footer_added_for_preview": preview_only,
        "automatic_sending_performed": False,
        "human_review_required": True,
        "validation_errors": errors,
    }
    return result, audit


HEADERS = [
    "优先级", "姓名", "当前机构", "论文角色", "赛道", "人才标签", "JD匹配分", "学术标签",
    "公开机构邮箱", "院系待转邮箱", "Homepage", "单版本触达素材",
    "地区", "竞品人才", "引用量", "H指数", "优势", "缺口", "异动窗口",
    "触达模式", "暂缓原因", "院系代转邮件", "A学术深耕邮件", "B全职流动邮件", "C兼职顾问邮件",
    "7天跟进", "LinkedIn附言", "LinkedIn私信", "证据URL", "检索时间", "合规状态",
]

ULTRA_FAST_HEADERS = [
    "优先级", "姓名", "当前机构", "论文角色", "研究赛道", "人才标签", "JD匹配分", "学术标签",
    "个人公开机构邮箱", "院系/机构待转邮箱", "Homepage", "单版本触达素材",
    "论文机构", "触达模式", "当前信息置信度", "证据URL",
]


def email_to_cell(message: dict[str, Any] | None) -> str:
    if not message:
        return ""
    return "\n\n".join(
        value for value in (
            f"Subject: {compact(message.get('subject'))}",
            compact(message.get("body")),
            compact(message.get("privacy_footer")),
        )
        if value
    )


def candidate_row(candidate: dict[str, Any], rendered: dict[str, Any], audit: dict[str, Any]) -> list[Any]:
    profile = candidate.get("profile") or {}
    policy = directive(candidate)
    initial = rendered.get("initial_emails") or {}
    public_emails = [
        item.get("email") for item in sendable_contacts(candidate, "public_institutional")
    ]
    department_emails = [
        item.get("email") for item in sendable_contacts(candidate, "departmental")
    ]
    strengths = (candidate.get("jd_match") or {}).get("strengths") or []
    gaps = (candidate.get("jd_match") or {}).get("gaps") or []
    hold_reasons = as_list(policy.get("hold_reasons")) + rendered.get("validation_errors", [])
    evidence = as_list(candidate.get("evidence"))
    evidence_urls = unique_strings(
        item.get("url") for item in evidence if isinstance(item, dict)
    )
    if not evidence_urls:
        anchor_url = compact(candidate_anchor(candidate).get("source_url"))
        evidence_urls = [anchor_url] if anchor_url else []
    raceway = first_value(profile.get("research_directions") or [])
    region = compact(profile.get("region") or policy.get("locale"))
    academic_tags = unique_strings(profile.get("academic_tags") or [])
    talent_tags = unique_strings(as_list(profile.get("talent_tags")) + academic_tags)
    homepage_urls = unique_strings(
        item.get("url") if isinstance(item, dict) else item
        for item in as_list(profile.get("homepages"))
    )
    quick_variant = compact(policy.get("quick_outreach_variant"))
    variant_map = {
        "academic": "A_academic_depth",
        "full_time": "B_full_time_mobility",
        "advisory": "C_advisory",
    }
    if quick_variant == "admin_forward":
        single_outreach = rendered.get("admin_forward")
    elif quick_variant == "archive_note":
        single_outreach = None
    else:
        single_outreach = initial.get(variant_map.get(quick_variant, ""))
    single_outreach_cell = (
        compact(rendered.get("internal_note"))
        if quick_variant == "archive_note"
        else email_to_cell(single_outreach)
    )
    return [
        compact(candidate.get("priority") or policy.get("priority")) or "未提供（待Agent判定）",
        candidate_name(candidate),
        "; ".join(unique_strings(profile.get("institutions") or [])),
        compact(candidate.get("paper_role") or profile.get("paper_role"))
        or "未提供（待核验）",
        raceway,
        "; ".join(talent_tags) or "未检索到（待核验）",
        (candidate.get("jd_match") or {}).get("score"),
        "; ".join(academic_tags) or "未检索到（待核验）",
        "; ".join(public_emails) or "未检索到（待核验）",
        "; ".join(department_emails) or "未检索到（待核验）",
        "; ".join(homepage_urls) or "未检索到（待核验）",
        single_outreach_cell or "未生成（不可发送/待补充）",
        region,
        "是" if policy.get("competitor_talent") is True else "否",
        get_metric(candidate, "citations"),
        get_metric(candidate, "h_index"),
        "; ".join(strengths),
        "; ".join(gaps),
        (candidate.get("mobility") or {}).get("window"),
        policy.get("mode"),
        "; ".join(unique_strings(hold_reasons)),
        email_to_cell(rendered.get("admin_forward")),
        email_to_cell(initial.get("A_academic_depth")),
        email_to_cell(initial.get("B_full_time_mobility")),
        email_to_cell(initial.get("C_advisory")),
        email_to_cell(rendered.get("follow_up")),
        compact((rendered.get("linkedin") or {}).get("connection_note")),
        compact((rendered.get("linkedin") or {}).get("direct_message")),
        "; ".join(evidence_urls),
        candidate.get("generated_at") or payload_generated_at(candidate),
        "可发送-待人工复核" if audit["sendable"] else "不可发送/仅存档",
    ]


def payload_generated_at(candidate: dict[str, Any]) -> str:
    return compact(candidate.get("generated_at")) or utc_now()


def ultra_fast_candidate_row(
    candidate: dict[str, Any],
    rendered: dict[str, Any],
    audit: dict[str, Any],
) -> list[Any]:
    """Return the fixed, self-contained 极速处理 row without cross-sheet formulas."""
    full_row = candidate_row(candidate, rendered, audit)
    profile = candidate.get("profile") or {}
    identity = profile.get("identity") or candidate.get("identity") or {}
    paper_institutions = unique_strings(
        as_list(candidate.get("paper_institutions"))
        + as_list(candidate.get("paper_institution"))
        + as_list(profile.get("paper_institutions"))
        + as_list(profile.get("paper_institution"))
    )
    confidence = (
        candidate.get("current_info_confidence")
        or profile.get("current_info_confidence")
        or identity.get("confidence")
        or identity.get("ambiguity")
        or "未提供（待核验）"
    )
    score = full_row[6]
    if score is None or compact(score) == "":
        score = "未评分（未提供JD）"
    return [
        *full_row[:6],
        score,
        *full_row[7:12],
        "; ".join(paper_institutions) or "未提供（待核验）",
        compact(directive(candidate).get("mode")) or "未提供（待Agent判定）",
        compact(confidence),
        full_row[28],
    ]


def write_ultra_fast_workbook(
    path: Path,
    candidates: list[dict[str, Any]],
    rendered: list[dict[str, Any]],
    audits: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "候选人清单"
    sheet.append(ULTRA_FAST_HEADERS)
    for candidate, message, audit in zip(candidates, rendered, audits):
        sheet.append([
            safe_cell(value)
            for value in ultra_fast_candidate_row(candidate, message, audit)
        ])

    header_fill = PatternFill("solid", fgColor="2D74A1")
    priority_fills = {
        "A": PatternFill("solid", fgColor="E7F4EC"),
        "B": PatternFill("solid", fgColor="FFF4D6"),
        "C": PatternFill("solid", fgColor="FBE5E3"),
    }
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        priority = compact(row[0].value)
        if priority in priority_fills:
            row[0].fill = priority_fills[priority]
            row[0].font = Font(bold=True)
            row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[6].alignment = Alignment(horizontal="center", vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = [10, 18, 30, 22, 34, 30, 12, 30, 30, 32, 32, 65, 34, 24, 30, 55]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_workbook(
    path: Path,
    candidates: list[dict[str, Any]],
    rendered: list[dict[str, Any]],
    audits: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人才清单"
    sheet.append(HEADERS)
    for candidate, message, audit in zip(candidates, rendered, audits):
        sheet.append([safe_cell(value) for value in candidate_row(candidate, message, audit)])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        1: 18, 2: 28, 3: 24, 4: 12, 5: 24, 6: 12, 7: 12, 8: 12,
        9: 12, 10: 28, 11: 28, 12: 28, 13: 28, 14: 24, 15: 12, 16: 16,
        17: 30, 18: 60, 19: 60, 20: 60, 21: 60, 22: 60, 23: 42, 24: 55,
        25: 55, 26: 24, 27: 20,
    }
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width

    audit_sheet = workbook.create_sheet("合规审计")
    audit_headers = [
        "candidate_key", "name", "requested_mode", "sendable", "preview_only",
        "recipient_source_url", "public_institutional", "gdpr_footer_added",
        "placeholder_footer_for_preview", "automatic_sending",
        "human_review_required", "validation_errors",
    ]
    audit_sheet.append(audit_headers)
    for audit in audits:
        audit_sheet.append(
            [
                safe_cell(audit.get("candidate_key")),
                safe_cell(audit.get("name")),
                safe_cell(audit.get("requested_mode")),
                audit.get("sendable"),
                audit.get("preview_only"),
                safe_cell(audit.get("recipient_source_url")),
                audit.get("recipient_is_public_institutional"),
                audit.get("gdpr_footer_added"),
                audit.get("placeholder_footer_added_for_preview"),
                audit.get("automatic_sending_performed"),
                audit.get("human_review_required"),
                safe_cell(audit.get("validation_errors")),
            ]
        )
    for cell in audit_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    audit_sheet.freeze_panes = "A2"
    audit_sheet.auto_filter.ref = audit_sheet.dimensions
    for column in range(1, len(audit_headers) + 1):
        audit_sheet.column_dimensions[get_column_letter(column)].width = 24

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Render localized recruiter outreach and export an audited Excel roster."
    )
    parser.add_argument("--input", required=True, help="Standardized batch profile JSON.")
    parser.add_argument("--output-dir", required=True, help="Output directory.")
    parser.add_argument("--xlsx-name", default="academic_talent_roster.xlsx")
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow replacing the three known output files.",
    )
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        payload = json.loads(Path(args.input).expanduser().read_text(encoding="utf-8"))
        candidates = payload.get("candidates")
        if not isinstance(candidates, list):
            raise ValueError("Input must contain a candidates list.")
        export_profile = compact(payload.get("export_profile"))
        if export_profile not in {"", ULTRA_FAST_EXPORT_PROFILE}:
            raise ValueError(f"Unsupported export_profile: {export_profile}")
        output_dir = Path(args.output_dir).expanduser().resolve()
        output_dir.mkdir(parents=True, exist_ok=True)
        message_path = output_dir / "messages.json"
        audit_path = output_dir / "compliance_audit.json"
        workbook_path = output_dir / args.xlsx_name
        targets = (
            [workbook_path]
            if export_profile == ULTRA_FAST_EXPORT_PROFILE
            else [message_path, audit_path, workbook_path]
        )
        existing = [path for path in targets if path.exists()]
        if existing and not args.overwrite:
            raise FileExistsError(
                "Refusing to overwrite existing outputs: "
                + ", ".join(str(path) for path in existing)
            )
        suppression = {
            compact(value).casefold()
            for value in as_list(payload.get("suppression_list"))
            if compact(value)
        }
        controller_errors = validate_controller(payload)
        rendered: list[dict[str, Any]] = []
        audits: list[dict[str, Any]] = []
        for candidate in candidates:
            message, audit = render_candidate(
                payload, candidate, suppression, controller_errors
            )
            rendered.append(message)
            audits.append(audit)
        messages_payload = {
            "schema_version": SCHEMA_VERSION,
            "run_id": payload.get("run_id"),
            "generated_at": utc_now(),
            "automatic_sending_performed": False,
            "human_review_required": True,
            "candidates": rendered,
        }
        audit_payload = {
            "schema_version": SCHEMA_VERSION,
            "run_id": payload.get("run_id"),
            "generated_at": utc_now(),
            "sendable_count": sum(item["sendable"] for item in audits),
            "draft_preview_count": sum(item["preview_only"] for item in audits),
            "blocked_or_archive_count": sum(not item["sendable"] for item in audits),
            "records": audits,
        }
        if export_profile == ULTRA_FAST_EXPORT_PROFILE:
            write_ultra_fast_workbook(workbook_path, candidates, rendered, audits)
        else:
            message_path.write_text(
                json.dumps(messages_payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            audit_path.write_text(
                json.dumps(audit_payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            write_workbook(workbook_path, candidates, rendered, audits)
        print(
            json.dumps(
                {
                    "messages": (
                        None if export_profile == ULTRA_FAST_EXPORT_PROFILE
                        else str(message_path)
                    ),
                    "audit": (
                        None if export_profile == ULTRA_FAST_EXPORT_PROFILE
                        else str(audit_path)
                    ),
                    "workbook": str(workbook_path),
                    "export_profile": export_profile or "full",
                    "candidates": len(candidates),
                    "sendable": audit_payload["sendable_count"],
                },
                ensure_ascii=False,
            )
        )
        return 0
    except Exception as exc:
        print(f"{type(exc).__name__}: {compact(str(exc))}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())

