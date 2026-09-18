"""Build a polished client-launch DOCX from structured Markdown.

Usage:
  python build_client_launch_docx.py input.md output.docx \
    --title "客户 岗位 客户启动与寻访手册" \
    --subtitle "客户理解 人才画像 电话筛选 寻访话术 执行计划" \
    --meta "适用对象=TTC 项目顾问和寻访顾问" \
    --meta "资料时点=2026 年 9 月 14 日" \
    --confidentiality "TTC 内部使用"
"""

from __future__ import annotations

import argparse
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
NAVY = "243B53"
PALE_BLUE = "EFF5FA"
PALE_GRAY = "F5F7F9"
BORDER = "D9D9D9"


def set_run_font(run, size=10.5, bold=None, color=BLACK, latin="Aptos", east_asia="Microsoft YaHei"):
    run.font.name = latin
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_borders(table, color=BORDER, size="6"):
    pr = table._tbl.tblPr
    borders = pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    shown = OxmlElement("w:t")
    shown.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, shown, end])
    paragraph.add_run(" 页")


def add_hyperlink(paragraph, text, url):
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F5D8A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.extend([fonts, color, underline])
    r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)")


def add_inline(paragraph, text, size=10.5, color=BLACK):
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos:match.start()]), size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            set_run_font(paragraph.add_run(token[2:-2]), size=size, bold=True, color=color)
        elif token.startswith("["):
            parsed = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if parsed:
                add_hyperlink(paragraph, parsed.group(1), parsed.group(2))
        else:
            set_run_font(
                paragraph.add_run(token[1:-1]),
                size=max(size - 0.5, 8.5),
                color=RGBColor(35, 35, 35),
                latin="Consolas",
            )
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), size=size, color=color)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.22

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(12)

    for name, size, before, after in (
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.first_line_indent = Inches(-0.16)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.18


def cover(doc, title, subtitle, metadata, confidentiality):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(46)
    set_run_font(p.add_run("TTC 人才寻访"), size=10, bold=True, color=RGBColor(36, 59, 83))

    p = doc.add_paragraph(style="Title")
    p.add_run(title)

    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(28)
        set_run_font(p.add_run(subtitle), size=11, color=RGBColor(72, 87, 101))

    items = list(metadata)
    if not any(k == "版本日期" for k, _ in items):
        items.append(("版本日期", date.today().isoformat()))
    for label, value in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run_font(p.add_run(f"{label}  "), size=10.5, bold=True, color=RGBColor(36, 59, 83))
        set_run_font(p.add_run(value), size=10.5)

    if confidentiality:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(40)
        set_run_font(p.add_run(f"保密等级  {confidentiality}"), size=9.5, bold=True, color=RGBColor(139, 49, 49))
    doc.add_page_break()


def parse_table(lines):
    return [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]


def add_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_borders(table)
    repeat_header(table.rows[0])
    for r_i, row in enumerate(rows):
        for c_i, value in enumerate(row):
            cell = table.cell(r_i, c_i)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            shade_cell(cell, NAVY if r_i == 0 else (PALE_BLUE if r_i % 2 == 0 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_i == 0 or len(value) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(
                p,
                value,
                size=9.2,
                color=RGBColor(255, 255, 255) if r_i == 0 else BLACK,
            )
            if r_i == 0:
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.05)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.04
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GRAY)
    p_pr.append(shd)
    for i, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, size=8.3, color=RGBColor(35, 35, 35), latin="Consolas")
        if i < len(lines) - 1:
            run.add_break()


def markdown_to_doc(doc, text):
    lines = text.splitlines()
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            if in_code:
                add_code(doc, code_lines)
                in_code = False
                code_lines = []
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(re.sub(r"^##\s+", "", line), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(re.sub(r"^###\s+", "", line), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(re.sub(r"^####\s+", "", line), level=3)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
            table_lines = [line]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_table(doc, parse_table(table_lines))
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        if line.startswith("> "):
            line = line[2:]

        parts = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt.strip() or nxt.startswith(("#", "- ", "|", "```", "> ")) or re.match(r"^\d+\.\s+", nxt):
                break
            parts.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(part.strip() for part in parts))


def parse_meta(values):
    items = []
    for value in values:
        if "=" not in value:
            raise SystemExit(f"--meta must be LABEL=VALUE: {value}")
        label, content = value.split("=", 1)
        items.append((label.strip(), content.strip()))
    return items


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--title", required=True)
    parser.add_argument("--subtitle", default="客户理解 人才画像 电话筛选 寻访话术 执行计划")
    parser.add_argument("--meta", action="append", default=[])
    parser.add_argument("--confidentiality", default="TTC 内部使用")
    args = parser.parse_args()

    doc = Document()
    setup(doc)
    cover(doc, args.title, args.subtitle, parse_meta(args.meta), args.confidentiality)
    markdown_to_doc(doc, args.input.read_text(encoding="utf-8"))
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        add_page_field(p)
        for run in p.runs:
            set_run_font(run, size=8.5, color=RGBColor(100, 110, 120))

    doc.core_properties.title = args.title
    doc.core_properties.subject = args.subtitle
    doc.core_properties.author = "TTC"
    doc.core_properties.comments = ""
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()

